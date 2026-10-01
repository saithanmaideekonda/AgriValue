"""Integration checks for the finalized AgriValue workbook and Flask app."""

from __future__ import annotations

import io
import json
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import app as agri_app
from utils.calculations import (
    InputError,
    calculate_financials,
    make_quantity_range,
    quantity_in_price_units,
    read_nonnegative_number,
    recalculate_returns,
    supported_quantity_units,
)
from utils.data_processing import (
    EXPECTED_COLUMNS,
    inspect_cleaning,
    load_source_data,
    make_dataset_summary,
    prepare_market_data,
)
from utils.model_training import FEATURES, TARGET
from utils.prediction import (
    commodities_for_analysis,
    districts_for_analysis,
    grades_for_analysis,
    markets_for_analysis,
    states_for_analysis,
    varieties_for_analysis,
)


class DatasetChecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = load_source_data()
        cls.cleaned = prepare_market_data(cls.source)

    def test_final_source_structure_and_dynamic_counts(self):
        summary = make_dataset_summary(self.source)
        self.assertEqual(self.source.shape, (10_272, 16))
        self.assertEqual(self.source.columns.tolist(), EXPECTED_COLUMNS)
        self.assertEqual(self.cleaned.columns.tolist(), EXPECTED_COLUMNS)
        self.assertEqual(summary["commodity_count"], 211)
        self.assertEqual(summary["state_count"], 17)
        self.assertEqual(summary["district_count"], 322)
        self.assertEqual(summary["market_count"], 1_118)
        self.assertIn("Telangana", self.source["State"].unique())
        self.assertEqual(summary["unique_dates"], 1)
        self.assertEqual(summary["valid_date_records"], 10_000)
        self.assertEqual(int(self.source["Arrival_Date"].isna().sum()), 272)
        self.assertEqual(int(self.source.duplicated().sum()), 27)
        self.assertEqual(len(commodities_for_analysis(self.cleaned)), self.source["Commodity"].nunique())
        self.assertEqual(summary["price_unit_distribution"], {"₹/quintal": 10_270, "₹/number": 2})

    def test_cleaning_keeps_master_columns_and_reports_source_issues(self):
        checks = inspect_cleaning(self.source)
        self.assertEqual(len(self.cleaned), 10_245)
        self.assertEqual(checks["source_records"], 10_272)
        self.assertEqual(self.source.shape[1], 16)
        self.assertEqual(checks["exact_duplicate_rows"], 27)
        self.assertEqual(checks["rows_removed"], 27)
        self.assertEqual(checks["missing_or_invalid_date_rows"], 272)
        self.assertEqual(checks["modal_price_iqr_outlier_rows_retained"], 509)
        self.assertEqual(self.cleaned.shape[1], 16)
        self.assertTrue((self.cleaned["Price_Unit"].astype(str).str.len() > 0).all())

    def test_master_data_is_not_changed_by_working_copy_cleaning(self):
        before = self.source.copy(deep=True)
        prepared = prepare_market_data(self.source)
        self.assertEqual(self.source.shape, before.shape)
        pd.testing.assert_frame_equal(self.source, before)
        self.assertEqual(prepared.columns.tolist(), EXPECTED_COLUMNS)

    def test_model_feature_list_uses_only_current_fields_and_excludes_target(self):
        self.assertEqual(TARGET, "Modal_Price")
        self.assertEqual(
            FEATURES,
            ["Commodity", "State", "District", "Market", "Variety", "Grade", "District_Region", "Price_Unit"],
        )
        self.assertNotIn(TARGET, FEATURES)
        self.assertEqual(set(FEATURES).difference(self.source.columns), set())


class CalculationChecks(unittest.TestCase):
    def test_recommendation_keeps_exact_ties_and_handles_no_valid_returns(self):
        rows = [
            {"market": "Market A", "estimated_price_source_units": 200, "quantity": 10,
             "transport_cost": 20, "expected_revenue": 2000, "estimated_net_return": 1980},
            {"market": "Market B", "estimated_price_source_units": 200, "quantity": 10,
             "transport_cost": 20, "expected_revenue": 2000, "estimated_net_return": 1980},
            {"market": "Unavailable", "estimated_price_source_units": 200, "quantity": 10,
             "transport_cost": 20, "expected_revenue": None, "estimated_net_return": None},
        ]
        tied = agri_app._recommendation_for_markets(rows)
        self.assertTrue(tied["available"])
        self.assertEqual([row["market"] for row in tied["markets"]], ["Market A", "Market B"])
        unavailable = agri_app._recommendation_for_markets([rows[2]])
        self.assertEqual(unavailable, {"available": False, "markets": []})

    def test_quintal_price_conversion_and_return(self):
        values = calculate_financials(2_000, "₹/quintal", 500, "kg", 1_200)
        self.assertEqual(values["expected_revenue"], 10_000)
        self.assertEqual(values["estimated_net_return"], 8_800)

    def test_number_price_uses_count_and_does_not_convert_weight(self):
        values = calculate_financials(6_400, "₹/number", 3, "number", 500)
        self.assertEqual(values["expected_revenue"], 19_200)
        self.assertEqual(values["estimated_net_return"], 18_700)
        with self.assertRaises(InputError):
            quantity_in_price_units(100, "kg", "₹/number")
        with self.assertRaises(InputError):
            calculate_financials(2_000, "₹/unknown", 100, "kg", 10)

    def test_quantity_range_uses_fixed_cost_and_marginal_difference(self):
        values = make_quantity_range(2_000, "₹/quintal", "kg", 100, 300, 100, 500)
        self.assertEqual([row["quantity"] for row in values], [100, 200, 300])
        self.assertEqual([row["expected_revenue"] for row in values], [2_000, 4_000, 6_000])
        self.assertEqual([row["estimated_net_return"] for row in values], [1_500, 3_500, 5_500])
        self.assertIsNone(values[0]["difference_from_previous"])
        self.assertEqual(values[1]["difference_from_previous"], 2_000)

    def test_calculation_input_validation(self):
        for value, label, allow_zero in [
            ("-1", "Quantity", False), ("NaN", "Transportation cost", True),
            ("inf", "Quantity", False), ("", "Quantity", False),
        ]:
            with self.subTest(value=value), self.assertRaises(InputError):
                read_nonnegative_number(value, label, allow_zero=allow_zero)
        with self.assertRaises(InputError):
            recalculate_returns([], 100, "kg", 0)
        self.assertEqual(supported_quantity_units("₹/quintal"), ["kg", "quintal"])
        self.assertEqual(supported_quantity_units("₹/number"), ["number"])


class FlaskRouteChecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        agri_app.refresh_runtime()
        agri_app.app.config.update(TESTING=True)
        cls.client = agri_app.app.test_client()
        cls.data = agri_app.MARKET_DATA
        cls.commodities = commodities_for_analysis(cls.data)
        cls.tomato = next(item for item in cls.commodities if item.casefold() == "tomato")
        cls.other = next(item for item in cls.commodities if item.casefold() != "tomato")
        cls.tomato_case = cls._combo(cls.tomato)
        cls.other_case = cls._combo(cls.other)
        cls.tel_case = cls._combo(cls.tomato, state="Telangana")

    @classmethod
    def _combo(cls, commodity, state=None):
        state_values = [state] if state else states_for_analysis(cls.data, commodity)
        for selected_state in state_values:
            for district in districts_for_analysis(cls.data, commodity, selected_state):
                markets = markets_for_analysis(cls.data, commodity, selected_state, district)
                if markets:
                    varieties = varieties_for_analysis(cls.data, commodity, selected_state, district, markets[:2])
                    variety = varieties[0] if varieties else ""
                    grades = grades_for_analysis(cls.data, commodity, selected_state, district, markets[:2], variety)
                    grade = grades[0] if grades else ""
                    selected_markets = markets[:2]
                    rows = cls.data[
                        cls.data["Commodity"].eq(commodity)
                        & cls.data["State"].eq(selected_state)
                        & cls.data["District"].eq(district)
                        & cls.data["Market"].isin(selected_markets)
                    ]
                    units = supported_quantity_units(rows["Price_Unit"].iloc[0])
                    quantity_unit = "kg" if "kg" in units else units[0]
                    return {
                        "commodity": commodity, "state": selected_state, "district": district,
                        "markets": selected_markets, "variety": variety, "grade": grade,
                        "quantity_unit": quantity_unit,
                    }
        raise AssertionError(f"No actual location records found for {commodity}")

    def _submit(self, case, **overrides):
        values = {**case, "quantity": "500", "transport_cost": "1200", **overrides}
        return self.client.post("/predict", data=values)

    def test_pages_start_and_dataset_loads(self):
        for path in ("/", "/analysis", "/results", "/visual-analysis", "/reports", "/model-info", "/dataset-summary", "/about"):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200, response.get_data(as_text=True)[:300])
        self.assertEqual(agri_app.RAW_DATA.shape, (10_272, 16))
        self.assertEqual(agri_app.MARKET_DATA.shape[1], 16)
        self.assertIsNotNone(agri_app.MODEL)
        self.assertIsNone(agri_app.DATASET_ERROR)
        self.assertIsNone(agri_app.MODEL_ERROR)

    def test_dynamic_state_to_grade_filtering(self):
        self.assertEqual(self.client.get("/api/states").status_code, 200)
        state = self.tel_case["state"]
        districts = self.client.get("/api/districts", query_string={"state": state}).get_json()["districts"]
        self.assertIn(self.tel_case["district"], districts)
        markets = self.client.get("/api/markets", query_string={
            "state": state, "district": self.tel_case["district"],
        }).get_json()["markets"]
        self.assertTrue(markets)
        commodity_options = self.client.get("/api/commodities", query_string={
            "state": state, "district": self.tel_case["district"], "market": self.tel_case["markets"][0],
        }).get_json()["commodities"]
        self.assertIn(self.tomato, commodity_options)
        tomato_markets = self.client.get("/api/markets", query_string={
            "state": state, "district": self.tel_case["district"], "commodity": self.tomato,
        }).get_json()["markets"]
        self.assertTrue(tomato_markets)
        query = {"state": state, "district": self.tel_case["district"], "commodity": self.tomato, "market": tomato_markets[0]}
        varieties = self.client.get("/api/varieties", query_string=query).get_json()["varieties"]
        if varieties:
            query["variety"] = varieties[0]
        grades = self.client.get("/api/grades", query_string=query).get_json()["grades"]
        self.assertIsInstance(grades, list)
        units = self.client.get("/api/quantity-units", query_string=query).get_json()
        self.assertIn("price_units", units)
        self.assertEqual(self.client.get("/api/states?commodity=Unknown").status_code, 200)

    def test_prediction_for_tomato_another_commodity_and_telanganan_records(self):
        for case in (self.tomato_case, self.other_case, self.tel_case):
            with self.subTest(commodity=case["commodity"], state=case["state"]):
                response = self._submit(case)
                self.assertEqual(response.status_code, 200, response.get_data(as_text=True)[:500])
                self.assertIn(b"Market Comparison", response.data)
                self.assertIn(b"Estimated Net Return", response.data)
                self.assertIn(case["commodity"].encode(), response.data)
                self.assertIn(case["markets"][0].encode(), response.data)

    def test_recommendation_matches_nine_market_comparison_and_updates_inputs(self):
        matching = self.data[
            self.data["State"].eq("Tamil Nadu")
            & self.data["District"].eq("Coimbatore")
            & self.data["Commodity"].str.casefold().eq("tomato")
            & self.data["Variety"].eq("Deshi")
            & self.data["Grade"].eq("Local")
        ]
        markets = matching["Market"].drop_duplicates().tolist()
        self.assertEqual(len(markets), 9)
        case = {"state": "Tamil Nadu", "district": "Coimbatore", "commodity": self.tomato,
                "markets": markets, "variety": "Deshi", "grade": "Local", "quantity_unit": "kg"}
        response = self._submit(case, quantity="500", transport_cost="1000")
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True)[:500])
        estimated = agri_app.estimate_markets(
            agri_app.MODEL, self.data, case["commodity"], case["state"], case["district"],
            markets, 500, "kg", 1000, variety="Deshi", grade="Local",
        )
        expected = agri_app._recommendation_for_markets(estimated)
        html = response.get_data(as_text=True)
        self.assertIn("Recommended Selling Market", html)
        self.assertIn("How this recommendation is calculated", html)
        self.assertEqual(html.count('data-recommendation-market'), len(expected["markets"]))
        for winner in expected["markets"]:
            self.assertIn(f'<h3 class="recommendation-market-name">{winner["market"]}</h3>', html)
            self.assertIn(f'{winner["estimated_net_return"]:,.2f}', html)

        with self.client.session_transaction() as sess:
            report_id = sess["latest_report_id"]
        recalculated = self.client.post("/calculate", json={
            "quantity": 1_000, "quantity_unit": "kg", "transport_cost": 2_000,
            "report_id": report_id,
            "markets": [{"market": item["market"],
                         "estimated_price_source_units": item["estimated_price_source_units"],
                         "price_unit": item["price_unit"]} for item in estimated],
        })
        self.assertEqual(recalculated.status_code, 200, recalculated.get_data(as_text=True))
        updated = [
            {**item, **values, "estimated_price_source_units": item["estimated_price_source_units"]}
            for item, values in zip(estimated, recalculated.get_json()["markets"])
        ]
        updated_expected = agri_app._recommendation_for_markets(updated)
        latest = self.client.get("/results")
        self.assertEqual(latest.status_code, 200)
        updated_html = latest.get_data(as_text=True)
        self.assertIn("1,000.00 kg", updated_html)
        self.assertIn("₹2,000.00", updated_html)
        self.assertEqual(updated_html.count('data-recommendation-market'), len(updated_expected["markets"]))
        for winner in updated_expected["markets"]:
            self.assertIn(f'<h3 class="recommendation-market-name">{winner["market"]}</h3>', updated_html)

        single_response = self._submit({**case, "markets": [markets[0]]})
        self.assertEqual(single_response.status_code, 200)
        self.assertEqual(single_response.get_data(as_text=True).count('data-recommendation-market'), 1)

    def test_results_show_unavailable_when_no_market_has_valid_return(self):
        case = self.tomato_case
        result = agri_app.estimate_markets(
            agri_app.MODEL, self.data, case["commodity"], case["state"], case["district"],
            case["markets"], 500, case["quantity_unit"], 1200,
            variety=case["variety"], grade=case["grade"],
        )[0]
        result["expected_revenue"] = None
        result["estimated_net_return"] = None
        with patch.object(agri_app, "estimate_markets", return_value=[result]):
            response = self._submit({**case, "markets": [case["markets"][0]]})
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn("Recommendation unavailable", html)
        self.assertNotIn('data-recommendation-market', html)

    def test_comparison_results_reconcile_using_the_recorded_unit(self):
        case = self.tomato_case
        response = self._submit(case)
        self.assertEqual(response.status_code, 200)
        markets = agri_app.estimate_markets(
            agri_app.MODEL, agri_app.MARKET_DATA, case["commodity"], case["state"], case["district"],
            case["markets"], 500, case["quantity_unit"], 1_200,
            variety=case["variety"], grade=case["grade"],
        )
        self.assertTrue(markets)
        for item in markets:
            factor = quantity_in_price_units(500, case["quantity_unit"], item["price_unit"])
            self.assertAlmostEqual(item["expected_revenue"], item["estimated_price_source_units"] * factor)
            self.assertAlmostEqual(item["estimated_net_return"], item["expected_revenue"] - 1_200)
        self.assertNotIn(b"Best Market", response.data)

    def test_invalid_selection_and_numeric_values_are_rejected(self):
        case = self.tomato_case
        form = {**case, "quantity": "500", "transport_cost": "1200"}
        invalid_cases = [
            {**form, "commodity": "Not in data"},
            {**form, "state": "Unknown state"},
            {**form, "district": "Unknown district"},
            {**form, "markets": ["Unknown market"]},
            {**form, "markets": []},
            {**form, "quantity": ""},
            {**form, "quantity": "-10"},
            {**form, "quantity": "NaN"},
            {**form, "quantity_unit": "unknown"},
            {**form, "transport_cost": "-1"},
            {**form, "transport_cost": "inf"},
        ]
        for payload in invalid_cases:
            with self.subTest(payload=payload):
                self.assertIn(self.client.post("/predict", data=payload).status_code, (400, 422))

    def test_calculation_api_and_quantity_change_leave_price_unchanged(self):
        markets = [{"market": "Sample", "estimated_price_source_units": 2_000, "price_unit": "₹/quintal"}]
        response = self.client.post("/calculate", json={
            "quantity": 500, "quantity_unit": "kg", "transport_cost": 1_200, "markets": markets,
        })
        self.assertEqual(response.status_code, 200)
        result = response.get_json()["markets"][0]
        self.assertEqual(result["expected_revenue"], 10_000)
        self.assertEqual(result["estimated_net_return"], 8_800)
        changed = self.client.post("/calculate", json={
            "quantity": 1_000, "quantity_unit": "kg", "transport_cost": 2_000, "markets": markets,
        }).get_json()["markets"][0]
        self.assertEqual(changed["expected_revenue"], 20_000)
        self.assertEqual(changed["estimated_net_return"], 18_000)
        invalid = self.client.post("/calculate", json={
            "quantity": 0, "quantity_unit": "kg", "transport_cost": 1_200, "markets": markets,
        })
        self.assertEqual(invalid.status_code, 400)
        unsupported = self.client.post("/calculate", json={
            "quantity": 100, "quantity_unit": "kg", "transport_cost": 0,
            "markets": [{"market": "Counted item", "estimated_price_source_units": 100, "price_unit": "₹/number"}],
        })
        self.assertEqual(unsupported.status_code, 400)

    def test_visual_analysis_quantity_and_exports_without_stability_section(self):
        case = self.tel_case
        query = {
            "state": case["state"], "district": case["district"], "commodity": case["commodity"],
            "market": case["markets"][0], "variety": case["variety"], "grade": case["grade"],
            "quantity_unit": case["quantity_unit"], "analyze": "1",
            "minimum_quantity": 100, "maximum_quantity": 300, "quantity_step": 100, "transport_cost": 50,
        }
        response = self.client.get("/visual-analysis", query_string=query)
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True)[:500])
        self.assertIn(b"Market Price Comparison", response.data)
        self.assertNotIn(b"Historical Price Stability", response.data)
        self.assertNotIn(b"Coefficient of Variation", response.data)
        self.assertNotIn(b"More observations are needed", response.data)
        self.assertIn(b"Quantity Range Difference Analysis", response.data)
        self.assertNotIn(b"Historical Price Trend", response.data)
        self.assertNotIn(b"Standard Deviation", response.data)
        self.assertNotIn(b"variation describes these observations", response.data)
        with self.client.session_transaction() as sess:
            report_id = sess["latest_report_id"]
        csv_download = self.client.get(f"/reports/{report_id}/csv")
        xlsx_download = self.client.get(f"/reports/{report_id}/xlsx")
        pdf_download = self.client.get(f"/reports/{report_id}/pdf")
        self.assertEqual(csv_download.status_code, 200)
        self.assertTrue(csv_download.data.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(xlsx_download.status_code, 200)
        workbook = load_workbook(io.BytesIO(xlsx_download.data), read_only=True)
        self.assertEqual(workbook.sheetnames, ["Input Details", "Market Results", "Stability Data", "Quantity Range Data"])
        self.assertEqual(pdf_download.status_code, 200)
        self.assertTrue(pdf_download.data.startswith(b"%PDF"))

    def test_dataset_summary_and_model_info_show_actual_values(self):
        summary = make_dataset_summary(load_source_data())
        page = self.client.get("/dataset-summary")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"10,272", page.data)
        self.assertIn(b"16 columns", page.data)
        for column in EXPECTED_COLUMNS:
            self.assertIn(column.encode(), page.data)
        self.assertIn(b"Telangana", page.data)
        self.assertNotIn(b"Historical Price Trend", page.data)
        model_page = self.client.get("/model-info")
        metadata = agri_app._metadata()
        self.assertEqual(model_page.status_code, 200)
        self.assertIn(b"Linear Regression (log-transformed target)", model_page.data)
        self.assertIn(f"{metadata['training_records']:,}".encode(), model_page.data)
        self.assertEqual(metadata["features_used"], FEATURES)
        self.assertIsNotNone(metadata["mae"])
        self.assertIsNotNone(metadata["rmse"])
        self.assertIsNotNone(metadata["r2"])


if __name__ == "__main__":
    unittest.main()

